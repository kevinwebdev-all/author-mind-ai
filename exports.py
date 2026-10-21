import io
import re

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    PageBreak,
)
from docx import Document
from docx.shared import Inches


class NovelExporter:
    """Export a novel to TXT, PDF, or DOCX."""

    @staticmethod
    def _safe_filename(title: str) -> str:
        name = re.sub(r'[\\/*?:"<>|]', "", title).strip()
        name = re.sub(r"\s+", "_", name)
        return (name or "novel")[:100]

    @staticmethod
    def _chapter_content(chapter: dict) -> str:
        return str(chapter.get("content") or "").strip()

    def export_txt(self, title: str, chapters: list[dict]) -> tuple[str, bytes]:
        output = io.StringIO()
        output.write(title.strip() + "\n")
        output.write("=" * min(len(title.strip()), 70) + "\n\n")

        for chapter in chapters:
            content = self._chapter_content(chapter)
            if not content:
                continue

            number = chapter.get("chapter_no", "?")
            chapter_title = chapter.get("title") or f"Chapter {number}"

            output.write(f"{chapter_title}\n\n")
            output.write(content + "\n\n")

        filename = self._safe_filename(title) + ".txt"
        return filename, output.getvalue().encode("utf-8")

    def export_docx(
        self, title: str, chapters: list[dict]
    ) -> tuple[str, bytes]:
        document = Document()
        document.add_heading(title.strip() or "Untitled Novel", level=0)

        for chapter in chapters:
            content = self._chapter_content(chapter)
            if not content:
                continue

            number = chapter.get("chapter_no", "?")
            chapter_title = chapter.get("title") or f"Chapter {number}"

            document.add_page_break()
            document.add_heading(chapter_title, level=1)

            for paragraph in content.split("\n\n"):
                paragraph = paragraph.strip()
                if paragraph:
                    document.add_paragraph(paragraph)

        buffer = io.BytesIO()
        document.save(buffer)

        filename = self._safe_filename(title) + ".docx"
        return filename, buffer.getvalue()

    def export_pdf(
        self, title: str, chapters: list[dict]
    ) -> tuple[str, bytes]:
        buffer = io.BytesIO()
        styles = getSampleStyleSheet()

        document = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=60,
            leftMargin=60,
            topMargin=60,
            bottomMargin=60,
            title=title.strip() or "Untitled Novel",
        )

        story = [
            Paragraph(
                title.strip() or "Untitled Novel",
                styles["Title"],
            ),
            Spacer(1, 24),
        ]

        for chapter in chapters:
            content = self._chapter_content(chapter)
            if not content:
                continue

            number = chapter.get("chapter_no", "?")
            chapter_title = chapter.get("title") or f"Chapter {number}"

            story.append(PageBreak())
            story.append(Paragraph(chapter_title, styles["Heading1"]))
            story.append(Spacer(1, 12))

            for paragraph in content.split("\n\n"):
                paragraph = paragraph.strip()
                if paragraph:
                    safe_text = (
                        paragraph.replace("&", "&amp;")
                        .replace("<", "&lt;")
                        .replace(">", "&gt;")
                    )
                    story.append(
                        Paragraph(safe_text.replace("\n", "<br/>"), styles["BodyText"])
                    )
                    story.append(Spacer(1, 8))

        document.build(story)

        filename = self._safe_filename(title) + ".pdf"
        return filename, buffer.getvalue()
          
